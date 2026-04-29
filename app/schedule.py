from __future__ import annotations

import re
import warnings
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Iterable

import requests
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


MONTHS = {
    "январ": 1,
    "феврал": 2,
    "март": 3,
    "апрел": 4,
    "май": 5,
    "мая": 5,
    "июн": 6,
    "июл": 7,
    "август": 8,
    "сентябр": 9,
    "октябр": 10,
    "ноябр": 11,
    "декабр": 12,
}

WEEKDAY_WORDS = (
    "пн",
    "вт",
    "ср",
    "чт",
    "пт",
    "сб",
    "вс",
    "понедельник",
    "вторник",
    "среда",
    "четверг",
    "пятница",
    "суббота",
    "воскресенье",
)


@dataclass(frozen=True)
class ScheduleEntry:
    day: date
    teacher: str
    lesson: str
    room: str
    sheet_name: str
    row_number: int


@dataclass(frozen=True)
class Schedule:
    entries: tuple[ScheduleEntry, ...]
    loaded_at: datetime

    @property
    def dates(self) -> list[date]:
        return sorted({entry.day for entry in self.entries})

    @property
    def teachers(self) -> list[str]:
        return sorted({entry.teacher for entry in self.entries}, key=sort_key_ru)

    def entries_for_date(self, target: date) -> list[ScheduleEntry]:
        return sorted(
            (entry for entry in self.entries if entry.day == target),
            key=lambda item: (sort_key_ru(item.teacher), lesson_sort_key(item.lesson), item.room),
        )

    def search_teacher(self, query: str) -> list[ScheduleEntry]:
        needle = normalize_search(query)
        if not needle:
            return []
        return sorted(
            (
                entry
                for entry in self.entries
                if needle in normalize_search(entry.teacher)
            ),
            key=lambda item: (item.day, sort_key_ru(item.teacher), lesson_sort_key(item.lesson)),
        )

    def next_dates(self, today: date, limit: int = 12) -> list[date]:
        future = [item for item in self.dates if item >= today]
        return future[:limit]


class GoogleSheetScheduleRepository:
    def __init__(self, spreadsheet_id: str, timeout: int = 30) -> None:
        self.spreadsheet_id = spreadsheet_id
        self.timeout = timeout

    @property
    def xlsx_url(self) -> str:
        return (
            "https://docs.google.com/spreadsheets/d/"
            f"{self.spreadsheet_id}/export?format=xlsx"
        )

    def load(self) -> Schedule:
        response = requests.get(self.xlsx_url, timeout=self.timeout)
        response.raise_for_status()
        content_type = response.headers.get("content-type", "")
        if "html" in content_type.lower():
            raise RuntimeError(
                "Google returned HTML instead of XLSX. Check sharing access for the table."
            )

        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="DrawingML support is incomplete.*",
                category=UserWarning,
            )
            workbook = load_workbook(BytesIO(response.content), data_only=True, read_only=False)
        entries = parse_workbook(workbook)
        return Schedule(entries=tuple(entries), loaded_at=datetime.now())


def parse_workbook(workbook) -> list[ScheduleEntry]:
    entries: list[ScheduleEntry] = []
    for worksheet in workbook.worksheets:
        if getattr(worksheet, "sheet_state", "visible") != "visible":
            continue
        if worksheet.title.strip().upper() == "ШАБЛОН":
            continue
        entries.extend(parse_worksheet(worksheet))
    return entries


def parse_worksheet(worksheet) -> list[ScheduleEntry]:
    entries: list[ScheduleEntry] = []
    month_year = detect_month_year(worksheet)
    max_row = worksheet.max_row or 0
    row = 1

    while row <= max_row:
        row_values = list(row_values_for(worksheet, row, max_col=8))
        current_date = extract_date(row_values, month_year)
        if current_date is None:
            row += 1
            continue

        lesson_row, lesson_map = find_lesson_row(worksheet, row)
        data_start = lesson_row + 1 if lesson_row else row + 1
        row = data_start

        while row <= max_row:
            values = list(row_values_for(worksheet, row, max_col=8))
            if extract_date(values, month_year) is not None:
                break

            teacher = normalize_teacher(value_to_text(values[0]) if values else "")
            if teacher and has_room_values(values, lesson_map):
                for col_index, lesson in lesson_map.items():
                    room = room_to_text(values[col_index - 1] if col_index - 1 < len(values) else "")
                    if room:
                        entries.append(
                            ScheduleEntry(
                                day=current_date,
                                teacher=teacher,
                                lesson=lesson,
                                room=room,
                                sheet_name=worksheet.title,
                                row_number=row,
                            )
                        )
            row += 1

    return entries


def row_values_for(worksheet, row: int, max_col: int) -> Iterable[object]:
    for col in range(1, max_col + 1):
        yield worksheet.cell(row=row, column=col).value


def find_lesson_row(worksheet, date_row: int) -> tuple[int | None, dict[int, str]]:
    for candidate in range(date_row + 1, min(date_row + 4, worksheet.max_row or date_row) + 1):
        lesson_map: dict[int, str] = {}
        for col in range(2, 8):
            lesson = lesson_to_text(worksheet.cell(row=candidate, column=col).value)
            if lesson:
                lesson_map[col] = lesson
        if len(lesson_map) >= 2 and all(is_lesson_number(item) for item in lesson_map.values()):
            return candidate, lesson_map

    return None, {col: str(col - 1) for col in range(2, 8)}


def has_room_values(values: list[object], lesson_map: dict[int, str]) -> bool:
    return any(
        room_to_text(values[col - 1] if col - 1 < len(values) else "")
        for col in lesson_map
    )


def detect_month_year(worksheet) -> tuple[int | None, int | None]:
    candidates: list[str] = [worksheet.title]
    for row in range(1, min(worksheet.max_row or 1, 6) + 1):
        for value in row_values_for(worksheet, row, max_col=8):
            text = value_to_text(value)
            if text:
                candidates.append(text)

    month: int | None = None
    year: int | None = None
    for text in candidates:
        lowered = text.lower()
        for marker, number in MONTHS.items():
            if marker in lowered:
                month = number
                break
        year_match = re.search(r"20\d{2}", text)
        if year_match:
            year = int(year_match.group(0))
        if month and year:
            return month, year
    return month, year


def extract_date(values: list[object], month_year: tuple[int | None, int | None]) -> date | None:
    for value in values:
        parsed = value_to_date(value)
        if parsed:
            return parsed

    joined = " ".join(value_to_text(value) for value in values if value_to_text(value))
    lowered = joined.lower()
    if not any(word in lowered for word in WEEKDAY_WORDS):
        return None

    month, year = month_year
    if not month or not year:
        return None

    day_match = re.search(r"(?<!\d)([0-3]?\d)(?!\d)", joined)
    if not day_match:
        return None
    day = int(day_match.group(1))
    try:
        return date(year, month, day)
    except ValueError:
        return None


def value_to_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 40000 <= float(value) <= 60000:
        try:
            return from_excel(float(value)).date()
        except (TypeError, ValueError):
            return None
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None


def value_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def lesson_to_text(value: object) -> str:
    text = value_to_text(value)
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def room_to_text(value: object) -> str:
    text = value_to_text(value)
    if not text:
        return ""
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"(?<!\d)(\d+)\.0(?!\d)", r"\1", text)
    return text.strip()


def normalize_teacher(name: str) -> str:
    cleaned = re.sub(r"\s+", " ", name).strip(" .")
    if not cleaned:
        return ""

    parts = [normalize_teacher_part(part) for part in cleaned.split("/")]
    return "/".join(part for part in parts if part)


def normalize_teacher_part(part: str) -> str:
    part = re.sub(r"\s+", " ", part).strip(" .")
    part = re.sub(r"\s*\.\s*", ".", part)
    match = re.match(r"^(.+?)\s+([А-ЯЁA-Z])\.?\s*([А-ЯЁA-Z])\.?$", part)
    if match:
        surname, first, second = match.groups()
        return f"{surname.strip()} {first}.{second}."
    return part


def normalize_search(text: str) -> str:
    return re.sub(r"[^0-9a-zа-яё]+", "", text.lower().replace("е\u0308", "ё"))


def sort_key_ru(text: str) -> str:
    return text.casefold().replace("ё", "е")


def lesson_sort_key(lesson: str) -> tuple[int, str]:
    if lesson.isdigit():
        return int(lesson), lesson
    return 999, lesson


def is_lesson_number(text: str) -> bool:
    return text.isdigit() and 1 <= int(text) <= 12


def group_entries_by_teacher(entries: Iterable[ScheduleEntry]) -> dict[str, list[ScheduleEntry]]:
    grouped: dict[str, list[ScheduleEntry]] = defaultdict(list)
    for entry in entries:
        grouped[entry.teacher].append(entry)
    return dict(grouped)
